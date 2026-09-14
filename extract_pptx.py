import sys
import collections 
import collections.abc
from pptx import Presentation

def extract_text(filepath, out_f):
    out_f.write(f"--- Extracting from {filepath} ---\n")
    try:
        prs = Presentation(filepath)
        for i, slide in enumerate(prs.slides):
            out_f.write(f"Slide {i+1}:\n")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    out_f.write(shape.text + "\n")
    except Exception as e:
        out_f.write(f"Error reading {filepath}: {e}\n")

if __name__ == "__main__":
    with open("output_pptx_fixed.txt", "w", encoding="utf-8") as out_f:
        for f in sys.argv[1:]:
            extract_text(f, out_f)
